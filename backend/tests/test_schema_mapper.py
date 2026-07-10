"""M5 Schema-Mapping Agent tests.

Three layers:
  * Unit: headers_match_canonical / map_columns (deterministic) / build_remap /
    remap_cells / validate_remap quality gate
  * LLM backend: use a fake model to test the sanitization logic (invalid fields cleared)
    and the fallback-to-deterministic on failure
  * End-to-end:
      - template "shifted right by one column", auto-aligned import
      - "rename + shift" combination, auto-aligned import
      - preview -> column_overrides -> import human-review roundtrip
      - heavy renaming -> SchemaMappingRejected

No API key needed (the LLM path uses a stub).
"""

from __future__ import annotations

import io
from pathlib import Path

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

from app.agents import schema_mapper as sm
from tests._db_fixture import setup_test_db

TEMPLATE = (
    Path(__file__).resolve().parent.parent.parent / "uploads" / "template" / "EcoTEA Endo WP1.xlsx"
)
HAS_TEMPLATE = TEMPLATE.exists()


# -----------------------------------------------------------------------------
# Unit tests
# -----------------------------------------------------------------------------
def test_standard_field_count():
    # A..AL = 38 standard fields, unique column positions
    assert len(sm.STANDARD_FIELDS) == 38
    cols = [f.column for f in sm.STANDARD_FIELDS]
    assert len(set(cols)) == 38


def test_deterministic_recovers_renamed_headers():
    blobs = {
        "B": "owner of data",  # data_owner
        "R": "build cost",  # capex
        "T": "FO&M charge",  # fixed_opex
        "O": "emissions factor",  # emission_factor
        "H": "asset name",  # technology_code
    }
    mapping = sm.map_columns(blobs, use_llm=False)
    got = {s.excel_column: s.target_field for s in mapping.suggestions}
    assert got["B"] == "data_owner"
    assert got["R"] == "capex"
    assert got["T"] == "fixed_opex"
    assert got["O"] == "emission_factor"
    assert got["H"] == "technology_code"


def test_one_field_one_column():
    # Both columns look like capex, but only one may be assigned capex
    blobs = {"R": "capex", "S": "capex build cost"}
    mapping = sm.map_columns(blobs, use_llm=False)
    fields = [s.target_field for s in mapping.suggestions if s.target_field == "capex"]
    assert len(fields) == 1


def test_build_remap_thresholds():
    mapping = sm.ColumnMapping(
        suggestions=[
            sm.ColumnSuggestion(excel_column="F", target_field="capex", confidence=0.95),
            sm.ColumnSuggestion(excel_column="G", target_field="fixed_opex", confidence=0.75),
            sm.ColumnSuggestion(excel_column="X", target_field="heat_rate", confidence=0.3),
        ]
    )
    remap, review = sm.build_remap(mapping)
    assert remap == {"F": "R"}  # high confidence -> moved to canonical column R
    assert [s.excel_column for s in review] == ["G"]  # medium confidence -> review
    # low-confidence X is dropped


def test_remap_cells_moves_values():
    remap = {"F": "R", "G": "T"}
    cells = {"F": 100, "G": 50, "Z": "ignored"}
    out = sm.remap_cells(cells, remap)
    assert out == {"R": 100, "T": 50}


def test_canonical_cells_from_raw_replays_remap():
    """Regression (#3): a remapped pending row's stored raw_cells (original + meta) must be
    reconstructed to canonical positions for conflict review / resolve."""
    from app.services.excel_importer import REMAP_META_KEY, _canonical_cells_from_raw

    # Standard layout (no meta) -> returned as-is.
    assert _canonical_cells_from_raw({"A": "Power", "R": 1572.0}) == {"A": "Power", "R": 1572.0}

    # Remapped pending row: original column A is blank (shifted file), the sector text is in B,
    # capex in S; the stashed remap B->A, S->R must be replayed and the meta key stripped.
    raw = {"A": None, "B": "Industry", "S": 1572.0, REMAP_META_KEY: {"B": "A", "S": "R"}}
    out = _canonical_cells_from_raw(raw)
    assert out == {"A": "Industry", "R": 1572.0}
    assert REMAP_META_KEY not in out

    # Defensive: non-dict input.
    assert _canonical_cells_from_raw(None) == {}


def test_build_remap_prefers_higher_confidence_on_duplicate_target():
    """The LLM backend may return two columns for the same standard field — the higher score must win, regardless of list order."""
    mapping = sm.ColumnMapping(
        suggestions=[
            sm.ColumnSuggestion(excel_column="C", target_field="capex", confidence=0.91),
            sm.ColumnSuggestion(excel_column="F", target_field="capex", confidence=0.98),
        ]
    )
    remap, _ = sm.build_remap(mapping)
    assert remap == {
        "F": "R"
    }, "column F at confidence 0.98 should win, not column C which appears first"


def test_validate_remap_rejects_missing_core_fields():
    """Core columns (technology_code/data_year) not aligned -> reject."""
    # Only capex aligned, core columns all missing
    remap = {"F": "R"}
    with pytest.raises(sm.SchemaMappingRejected) as exc_info:
        sm.validate_remap(remap, [], sheet_name="Power")
    msg = str(exc_info.value)
    assert "technology_code" in msg
    assert "preview" in msg  # the error message should point the user to human review


def test_validate_remap_rejects_low_coverage():
    """Core columns present but overall coverage below the threshold -> reject."""
    remap = {
        "H": sm.FIELD_BY_NAME["technology_code"].column,
        "K": sm.FIELD_BY_NAME["data_year"].column,
    }
    assert len(remap) / len(sm.STANDARD_FIELDS) < sm.MIN_AUTO_MAPPED_RATIO
    with pytest.raises(sm.SchemaMappingRejected, match="coverage"):
        sm.validate_remap(remap, [], sheet_name="Power")


def test_validate_remap_passes_good_mapping():
    """Core columns present + coverage met -> allowed."""
    remap = {spec.column: spec.column for spec in sm.STANDARD_FIELDS}
    sm.validate_remap(remap, [], sheet_name="Power")  # should not raise


# -----------------------------------------------------------------------------
# LLM backend (fake model stub, no API key needed)
# -----------------------------------------------------------------------------
class _FakeStructuredLLM:
    """Stub for get_chat_model(): after with_structured_output, invoke returns a preset result."""

    def __init__(self, result: sm.ColumnMapping):
        self._result = result
        self.invoked_with: list | None = None

    def with_structured_output(self, schema):  # noqa: ANN001
        assert schema is sm.ColumnMapping
        return self

    def invoke(self, messages):  # noqa: ANN001
        self.invoked_with = messages
        return self._result


def test_llm_backend_cleans_invalid_target_field(monkeypatch):
    """The LLM hallucinates a nonexistent field name -> cleared and confidence zeroed."""
    fake = _FakeStructuredLLM(
        sm.ColumnMapping(
            suggestions=[
                sm.ColumnSuggestion(
                    excel_column="A", target_field="hallucinated_field", confidence=0.99
                ),
                sm.ColumnSuggestion(excel_column="B", target_field="capex", confidence=0.95),
            ]
        )
    )
    import app.llm.provider as provider

    monkeypatch.setattr(provider, "get_chat_model", lambda **kw: fake)

    mapping = sm.map_columns_llm({"A": "mystery", "B": "build cost"})
    by_col = {s.excel_column: s for s in mapping.suggestions}
    assert by_col["A"].target_field is None
    assert by_col["A"].confidence == 0.0
    assert by_col["B"].target_field == "capex"
    # The prompt should carry the few-shot examples and all standard fields
    system_text = fake.invoked_with[0].content
    assert "owner of data" in system_text  # few-shot example
    assert "technology_code" in system_text


def test_llm_failure_falls_back_to_deterministic(monkeypatch):
    """The LLM backend raises -> map_columns falls back to the deterministic backend automatically."""
    import app.llm.provider as provider

    def _boom(**kw):  # noqa: ANN003
        raise RuntimeError("no api key")

    monkeypatch.setattr(provider, "get_chat_model", _boom)
    mapping = sm.map_columns({"R": "build cost"}, use_llm=True)
    got = {s.excel_column: s.target_field for s in mapping.suggestions}
    assert got["R"] == "capex"


@pytest.mark.skipif(not HAS_TEMPLATE, reason="template file missing")
def test_template_is_canonical():
    wb = load_workbook(TEMPLATE, data_only=True, read_only=True)
    blobs = sm.extract_header_blobs(wb["Power"])
    wb.close()
    assert sm.headers_match_canonical(blobs) is True


# -----------------------------------------------------------------------------
# End-to-end: shift right by one column -> M5 auto-align -> import
# -----------------------------------------------------------------------------
def _shift_template_right(only_sheet: str = "Power") -> bytes:
    """Shift one sheet of the template right by one column (insert a blank column at the front), delete other sheets.

    After shifting, the data of canonical column R lands in S; the mapper must realign S back to R.
    """
    wb = load_workbook(TEMPLATE, data_only=True)
    for name in list(wb.sheetnames):
        if name != only_sheet:
            del wb[name]
    wb[only_sheet].insert_cols(1)  # shift all columns right by one; A becomes empty
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.skipif(not HAS_TEMPLATE, reason="template file missing")
def test_end_to_end_shifted_import_aligns_columns():
    from app.database import SessionLocal
    from app.services.excel_importer import import_excel

    setup_test_db()

    shifted = _shift_template_right("Power")

    # After shifting, the header is no longer the canonical layout
    wb = load_workbook(io.BytesIO(shifted), data_only=True, read_only=True)
    blobs = sm.extract_header_blobs(wb["Power"])
    wb.close()
    assert sm.headers_match_canonical(blobs) is False

    db = SessionLocal()
    try:
        result = import_excel(
            db,
            file_bytes=shifted,
            file_name="shifted.xlsx",
            note="m5-shift-test",
            auto_map_columns=True,
            use_llm_mapping=False,  # deterministic backend
        )
        assert result.rows_imported > 0, "the shifted file should import rows after M5 alignment"

        # Assert capex actually landed in the canonical column (ecotea_parameter.capex non-null and positive)
        capex_rows = db.execute(
            text("SELECT COUNT(*) FROM technology_year_ecotea_parameter WHERE capex > 0")
        ).scalar()
        tech_codes = db.execute(
            text(
                "SELECT COUNT(*) FROM technology_process WHERE technology_code IS NOT NULL "
                "AND technology_code <> ''"
            )
        ).scalar()
    finally:
        db.close()

    assert capex_rows > 0, "capex should be relocated to the canonical column and written"
    assert tech_codes > 0, "technology_code (column H) should be aligned correctly"


# Common renames (simulating a template version bump), kept consistent with verify_schema_mapping.py
RENAME = {
    "data owner": "owner of data",
    "capex": "build cost",
    "fixed opex": "FO&M charge",
    "variable opex": "VO&M",
    "ef": "emissions factor",
    "technology/process": "asset name",
    "year of data": "year",
}


def _rename_and_shift_template(only_sheet: str = "Power") -> bytes:
    """Rename (row 2 synonyms) + shift right by one column: the worst common version-bump combination."""
    wb = load_workbook(TEMPLATE, data_only=True)
    for name in list(wb.sheetnames):
        if name != only_sheet:
            del wb[name]
    ws = wb[only_sheet]
    for cell in ws[2]:
        if cell.value and str(cell.value).lower() in RENAME:
            cell.value = RENAME[str(cell.value).lower()]
    ws.insert_cols(1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.mark.skipif(not HAS_TEMPLATE, reason="template file missing")
def test_end_to_end_renamed_and_shifted_import():
    """Rename + shift combination: after M5 alignment, capex / technology code still land in canonical columns."""
    from app.database import SessionLocal
    from app.services.excel_importer import import_excel

    setup_test_db()
    db = SessionLocal()
    try:
        result = import_excel(
            db,
            file_bytes=_rename_and_shift_template("Power"),
            file_name="renamed_shifted.xlsx",
            note="m5-rename-shift-test",
            auto_map_columns=True,
            use_llm_mapping=False,
        )
        assert result.rows_imported > 0
        # ImportResult must surface column-alignment warnings (auto-alignment was enabled)
        assert (
            result.column_warnings
        ), "there should be column-level warnings when Schema-Mapping is enabled"
        assert any("Auto-aligned" in w for w in result.column_warnings)

        capex_rows = db.execute(
            text("SELECT COUNT(*) FROM technology_year_ecotea_parameter WHERE capex > 0")
        ).scalar()
    finally:
        db.close()
    assert (
        capex_rows > 0
    ), "capex should still be relocated to the canonical column after rename+shift"


@pytest.mark.skipif(not HAS_TEMPLATE, reason="template file missing")
def test_preview_overrides_import_roundtrip():
    """Human-review roundtrip: preview suggests -> assemble column_overrides -> import applies them."""
    from app.database import SessionLocal
    from app.services.excel_importer import import_excel, preview_excel

    setup_test_db()
    shifted = _shift_template_right("Power")

    preview = preview_excel(file_bytes=shifted, file_name="shifted.xlsx")
    power = next(s for s in preview.sheets if s.sheet_name == "Power")
    assert power.column_mapping is not None
    assert power.column_mapping.layout_is_standard is False
    assert (
        preview.standard_fields
    ), "preview should return the standard field list for the frontend dropdown"

    # Simulate the user accepting all suggestions in ColumnMappingReview (confirm both auto + review)
    overrides = {
        s.excel_column: s.target_column
        for s in power.column_mapping.suggestions
        if s.target_column and s.status in ("auto", "review")
    }
    assert overrides, "the shifted file should have confirmable column suggestions"

    db = SessionLocal()
    try:
        result = import_excel(
            db,
            file_bytes=shifted,
            file_name="shifted.xlsx",
            note="m5-overrides-roundtrip",
            column_overrides={"Power": overrides},
            auto_map_columns=False,  # with overrides, don't call the Agent again
        )
        assert result.rows_imported > 0
        capex_rows = db.execute(
            text("SELECT COUNT(*) FROM technology_year_ecotea_parameter WHERE capex > 0")
        ).scalar()
    finally:
        db.close()
    assert capex_rows > 0, "user-confirmed overrides should be applied as-is"


@pytest.mark.skipif(not HAS_TEMPLATE, reason="template file missing")
def test_end_to_end_garbled_headers_rejected():
    """Heavy renaming (headers unrecognizable) -> reject the import with a hint, instead of silently producing broken data."""
    from app.database import SessionLocal
    from app.services.excel_importer import import_excel

    setup_test_db()
    wb = load_workbook(TEMPLATE, data_only=True)
    for name in list(wb.sheetnames):
        if name != "Power":
            del wb[name]
    ws = wb["Power"]
    # Replace all headers in rows 1-9 with meaningless placeholders, destroying all core-column signal
    for row in ws.iter_rows(min_row=1, max_row=9):
        for cell in row:
            if cell.value is not None:
                cell.value = f"mystery_{cell.column_letter}_{cell.row}"
    buf = io.BytesIO()
    wb.save(buf)

    db = SessionLocal()
    try:
        with pytest.raises(sm.SchemaMappingRejected, match="technology_code"):
            import_excel(
                db,
                file_bytes=buf.getvalue(),
                file_name="garbled.xlsx",
                auto_map_columns=True,
                use_llm_mapping=False,
            )
    finally:
        db.close()


@pytest.mark.skipif(not HAS_TEMPLATE, reason="template file missing")
def test_standard_template_still_imports_fast_path():
    """Regression: the standard template still takes the fast path (column_remap=None); import unaffected by M5."""
    from app.database import SessionLocal
    from app.services.excel_importer import import_excel

    setup_test_db()
    wb = load_workbook(TEMPLATE, data_only=True)
    for name in list(wb.sheetnames):
        if name != "Power":
            del wb[name]
    buf = io.BytesIO()
    wb.save(buf)

    db = SessionLocal()
    try:
        result = import_excel(
            db,
            file_bytes=buf.getvalue(),
            file_name="std.xlsx",
            auto_map_columns=True,
            use_llm_mapping=False,
        )
        assert result.rows_imported > 0
    finally:
        db.close()
